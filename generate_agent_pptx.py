#!/usr/bin/env python3
"""Build ESA_Telemanom_Agent_Supervisor.pptx from the approved slide outline.

  conda activate esa-agent
  pip install python-pptx
  python generate_agent_pptx.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

OUT = Path(__file__).resolve().parent / "ESA_Telemanom_Agent_Supervisor.pptx"

# Teal / slate palette (readable, not purple)
NAVY = RGBColor(0x0A, 0x3D, 0x40)
TEAL = RGBColor(0x0D, 0x73, 0x77)
TEAL_LIGHT = RGBColor(0xE6, 0xF3, 0xF2)
SLATE = RGBColor(0x14, 0x23, 0x2B)
MUTED = RGBColor(0x3D, 0x55, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xB7, 0xC8, 0xCC)
CARD = RGBColor(0xF3, 0xF6, 0xF7)


def _set_run(run, text: str, size: int, bold: bool = False, color=SLATE) -> None:
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _bar(slide, prs) -> None:
    """Top accent bar on content slides."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.08)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = TEAL
    shape.line.fill.background()


def _bullets(slide, items: list[str], top: float = 1.05, size: int = 18) -> None:
    box = _add_textbox(slide, Inches(0.6), Inches(top), Inches(12.1), Inches(5.8))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        text = item
        if text.startswith("  "):
            p.level = 1
            text = text.strip()
            if text.startswith("- "):
                text = text[2:]
        p.text = text
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.color.rgb = SLATE
            run.font.name = "Calibri"


def _table(slide, rows: list[list[str]], left, top, width, row_h=0.42) -> None:
    n_rows, n_cols = len(rows), len(rows[0])
    table_shape = slide.shapes.add_table(
        n_rows, n_cols, left, top, width, Inches(row_h * n_rows)
    )
    table = table_shape.table
    for r in range(n_rows):
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.text = rows[r][c]
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(14 if r > 0 else 15)
                    run.font.bold = r == 0
                    run.font.name = "Calibri"
                    run.font.color.rgb = WHITE if r == 0 else SLATE
            cell.fill.solid()
            cell.fill.fore_color.rgb = TEAL if r == 0 else (TEAL_LIGHT if r % 2 else WHITE)


def add_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # Background wash via full rect
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0xEE, 0xF4, 0xF5)
    bg.line.fill.background()

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.22), prs.slide_height
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()

    box = _add_textbox(slide, Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.2))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "ESA Telemanom Agent", 40, bold=True, color=NAVY)

    sub = _add_textbox(slide, Inches(0.9), Inches(3.4), Inches(11.5), Inches(0.6))
    p = sub.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(
        run,
        "Single tool-calling agent over RAG + Telemanom scoring",
        22,
        color=MUTED,
    )

    meta = _add_textbox(slide, Inches(0.9), Inches(4.4), Inches(11.5), Inches(1.0))
    tf = meta.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "Workspace:  esa-telemanom-agent", 16, color=SLATE)
    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    _set_run(run2, "Stack:  FAISS RAG  ·  Groq LLM  ·  Telemanom scoring  ·  Streamlit UI", 16, color=SLATE)


def add_content_slide(prs: Presentation, title: str) -> object:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    _bar(slide, prs)
    box = _add_textbox(slide, Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.6))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(run, title, 26, bold=True, color=NAVY)
    return slide


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # --- Title ---
    add_title_slide(prs)

    # --- 1 Context ---
    s = add_content_slide(prs, "1. Context — ESA Telemanom")
    _bullets(
        s,
        [
            "Telemanom-ESA on ESA-Mission1 telemetry (84_months)",
            "LSTM predicts the next window of target channels from recent inputs",
            "Large prediction error → dynamic thresholds → alarms",
            "We reuse the lab model pack as frozen artifacts:",
            "  - model.h5 + normalization (means / stds)",
            "  - contract.json (channel order, window, AD parameters)",
            "87 input channels → 58 target channels",
            "Look-back window: 250 timesteps",
            "This workspace does not retrain the model; it wraps it for an agent.",
        ],
        size=17,
    )

    # --- 2 Goal ---
    s = add_content_slide(prs, "2. Goal")
    note = _add_textbox(s, Inches(0.6), Inches(0.95), Inches(12), Inches(0.4))
    run = note.text_frame.paragraphs[0].add_run()
    _set_run(run, "Build a clear agent system that can:", 16, color=MUTED)
    _table(
        s,
        [
            ["Need", "How"],
            ["Answer questions about setup, channels, model facts", "RAG over local knowledge/"],
            ["Run anomaly scoring on a telemetry CSV slice", "Tool calling Telemanom inference"],
            ["Check alarms against ground-truth labels", "Tool for interval overlap (true / wrong / miss)"],
            ["Demonstrate the full flow", "Streamlit UI"],
        ],
        Inches(0.55),
        Inches(1.4),
        Inches(12.2),
        row_h=0.48,
    )
    foot = _add_textbox(s, Inches(0.6), Inches(4.2), Inches(12), Inches(1.2))
    tf = foot.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(
        run,
        "Approach: learn step by step — RAG → agent → scoring → combined tools + UI.",
        16,
        color=SLATE,
    )
    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    _set_run(run2, "One conda environment: esa-agent.", 16, color=SLATE)

    # --- 3 Architecture ---
    s = add_content_slide(prs, "3. System architecture")
    # Simple boxes instead of ASCII art
    def _card(slide, x, y, w, h, text, fill=TEAL_LIGHT, title_color=NAVY):
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        sh.line.color.rgb = TEAL
        tf = sh.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        run = tf.paragraphs[0].add_run()
        _set_run(run, text, 14, bold=True, color=title_color)

    _card(s, 3.8, 1.1, 5.5, 0.7, "User question (+ optional CSV slice)", WHITE, MUTED)
    _card(s, 3.8, 2.1, 5.5, 0.85, "Single LLM agent\n(Groq tool-calling)", TEAL, WHITE)
    _card(s, 0.5, 3.4, 3.7, 1.1, "retrieve_knowledge\n→ FAISS + knowledge/", TEAL_LIGHT)
    _card(s, 4.7, 3.4, 3.7, 1.1, "score_telemetry\n→ model/ + scoring/", TEAL_LIGHT)
    _card(s, 8.9, 3.4, 3.8, 1.1, "compare_alarms_to_labels\n→ alarms vs is_anomaly_*", TEAL_LIGHT)
    foot = _add_textbox(s, Inches(0.6), Inches(5.0), Inches(12), Inches(1.5))
    tf = foot.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "Mental model: the LLM chooses tools and narrates results.", 16, bold=True, color=NAVY)
    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    _set_run(
        run2,
        "Anomaly detection stays in Telemanom. Facts stay in RAG. Labels provide an audit check.",
        15,
        color=SLATE,
    )

    # --- 4 Knowledge + RAG ---
    s = add_content_slide(prs, "4. Knowledge base + RAG (FAISS)")
    left_t = _add_textbox(s, Inches(0.55), Inches(0.95), Inches(6), Inches(0.35))
    run = left_t.text_frame.paragraphs[0].add_run()
    _set_run(run, "Knowledge corpus", 16, bold=True, color=TEAL)
    _table(
        s,
        [
            ["Source", "Content"],
            ["channels.csv", "Channel metadata"],
            ["anomaly_types.csv", "Anomaly-type catalogue"],
            ["model_card.txt", "Architecture facts (window, LSTM, I/O)"],
            ["lab_results_84months.csv", "Compact lab-run summary for retrieval"],
        ],
        Inches(0.55),
        Inches(1.35),
        Inches(6.2),
        row_h=0.4,
    )
    right = _add_textbox(s, Inches(7.1), Inches(0.95), Inches(5.5), Inches(5.5))
    tf = right.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "RAG pipeline", 16, bold=True, color=TEAL)
    for line in [
        "1. Chunk the knowledge files",
        "2. Embed with MiniLM (all-MiniLM-L6-v2, local)",
        "3. Index with FAISS",
        "4. Query → top-k chunks (capped at 8)",
        "",
        "RAG answers documentation / setup questions.",
        "It does not store raw telemetry timesteps.",
    ]:
        p = tf.add_paragraph()
        p.space_after = Pt(6)
        run = p.add_run()
        _set_run(run, line, 15, color=SLATE)

    # --- 5 Build path ---
    s = add_content_slide(prs, "5. Build path (Step 1 → Step 4)")
    _table(
        s,
        [
            ["Step", "What we added"],
            ["1", "FAISS index: build + query"],
            ["2", "Single agent + retrieve_knowledge"],
            ["3", "Telemanom scoring library + CLI"],
            ["4", "Agent with 3 tools + Streamlit UI"],
        ],
        Inches(1.5),
        Inches(1.5),
        Inches(10),
        row_h=0.55,
    )
    foot = _add_textbox(s, Inches(0.6), Inches(4.8), Inches(12), Inches(1.2))
    tf = foot.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(
        run,
        "Each step remains runnable on its own for understanding and debugging.",
        16,
        color=SLATE,
    )
    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    _set_run(run2, "Later: multi-agent extension.", 16, color=SLATE)

    # --- 6 Single agent ---
    s = add_content_slide(prs, "6. Single agent")
    _bullets(
        s,
        [
            "One ReAct-style tool-calling loop",
            "LLM via Groq (API key in .env)",
            "Flow: question → optional tool calls → observations → short final answer",
            "Grounding rules:",
            "  - Answer from tool results",
            "  - Do not invent physical root causes",
            "  - Keep answers concise",
            "One agent, several tools — not a multi-agent system yet.",
        ],
        size=18,
    )

    # --- 7 Three tools ---
    s = add_content_slide(prs, "7. Three tools")
    _table(
        s,
        [
            ["Tool", "Role"],
            ["retrieve_knowledge", "Search FAISS over knowledge/; return top chunks"],
            [
                "score_telemetry",
                "Run Telemanom on a CSV slice; write scores/alarms; short summary",
            ],
            [
                "compare_alarms_to_labels",
                "Compare alarms to is_anomaly_* columns on the same slice",
            ],
        ],
        Inches(0.55),
        Inches(1.5),
        Inches(12.2),
        row_h=0.7,
    )
    foot = _add_textbox(s, Inches(0.6), Inches(5.0), Inches(12), Inches(0.6))
    run = foot.text_frame.paragraphs[0].add_run()
    _set_run(run, "The LLM picks which tool(s) to call from the user question.", 16, color=MUTED)

    # --- 8 ESA metrics ---
    s = add_content_slide(prs, "8. Evaluation metrics (ESA-ADB)")
    note = _add_textbox(s, Inches(0.6), Inches(0.95), Inches(12), Inches(0.45))
    run = note.text_frame.paragraphs[0].add_run()
    _set_run(
        run,
        "Same TimeEval ESAScores path as the ESA-ADB / Telemanom lab evaluation.",
        15,
        color=MUTED,
    )
    _table(
        s,
        [
            ["Metric family", "What it covers"],
            ["alarming_precision", "Extra / redundant detections"],
            ["EW_*", "Corrected event-wise precision, recall, F0.5"],
            ["AFF_*", "Affiliation-based precision, recall, F0.5"],
        ],
        Inches(0.55),
        Inches(1.5),
        Inches(12.2),
        row_h=0.55,
    )
    _bullets(
        s,
        [
            "Ground truth from ESA labels.csv (+ anomaly categories)",
            "Global prediction: alarm if any target channel fires (lab rule)",
            "Definitions match the lab Telemanom evaluation path",
        ],
        top=4.3,
        size=16,
    )

    # --- 9 Streamlit ---
    s = add_content_slide(prs, "9. Streamlit UI")
    _bullets(
        s,
        [
            "Question box and example prompts",
            "Sidebar: CSV path, skip_rows, max_rows, RAG k",
            "Shows the final answer and expandable tool steps (name, arguments, observation)",
            "Same agent backend as the CLI",
            "",
            "Run:",
            "  conda activate esa-agent",
            "  cd esa-telemanom-agent/step4_agent",
            "  streamlit run app.py",
        ],
        size=17,
    )

    # --- 10 Local vs LLM ---
    s = add_content_slide(prs, "10. What reaches the LLM vs what stays local")
    _table(
        s,
        [
            ["Goes to the LLM", "Stays local"],
            ["Retrieved knowledge chunks", "FAISS index + embedder"],
            [
                "Short score / compare summaries",
                "Full telemetry CSV, score/alarm files, model.h5",
            ],
        ],
        Inches(0.8),
        Inches(1.8),
        Inches(11.5),
        row_h=0.65,
    )
    foot = _add_textbox(s, Inches(0.6), Inches(4.5), Inches(12), Inches(1.2))
    tf = foot.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "Embeddings and Telemanom inference run on the machine.", 16, color=SLATE)
    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    _set_run(run2, "Only the chat completion call goes to Groq.", 16, color=SLATE)

    # --- 11 Scope ---
    s = add_content_slide(prs, "11. Scope")
    left = _add_textbox(s, Inches(0.55), Inches(1.1), Inches(5.8), Inches(5))
    tf = left.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "Can today", 18, bold=True, color=TEAL)
    for line in [
        "Grounded Q&A from project knowledge",
        "Score a telemetry CSV slice and summarize alarms",
        "Audit alarms against ground-truth labels on that slice",
        "Show tool traces in the UI",
    ]:
        p = tf.add_paragraph()
        p.space_before = Pt(10)
        run = p.add_run()
        _set_run(run, "•  " + line, 15, color=SLATE)

    right = _add_textbox(s, Inches(6.8), Inches(1.1), Inches(5.8), Inches(5))
    tf = right.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "Cannot today (by design)", 18, bold=True, color=TEAL)
    for line in [
        "Invent physical / subsystem root-cause explanations",
        "Replace a full-mission offline evaluation report",
        "Multi-agent specialist workflow (planned next)",
    ]:
        p = tf.add_paragraph()
        p.space_before = Pt(10)
        run = p.add_run()
        _set_run(run, "•  " + line, 15, color=SLATE)

    # --- 12 Summary ---
    s = add_content_slide(prs, "12. Summary")
    _table(
        s,
        [
            ["Piece", "Choice"],
            ["Retrieval", "FAISS + MiniLM over knowledge/"],
            ["Reasoning", "Single Groq tool-calling agent"],
            ["Detection", "Lab Telemanom model via scoring/"],
            ["Audit", "ESA-ADB ESAScores (EW / AFF / alarming)"],
            ["Interface", "Streamlit + CLI"],
        ],
        Inches(1.2),
        Inches(1.3),
        Inches(10.8),
        row_h=0.48,
    )
    foot = _add_textbox(s, Inches(0.7), Inches(5.0), Inches(11.8), Inches(1.3))
    tf = foot.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, "One line:", 14, bold=True, color=TEAL)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    run2 = p2.add_run()
    _set_run(
        run2,
        "A step-by-step single-agent system that retrieves ESA/Telemanom knowledge, "
        "runs the lab model on demand, and compares alarms to labels.",
        16,
        color=SLATE,
    )

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
